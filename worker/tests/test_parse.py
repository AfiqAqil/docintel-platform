"""Tests for deterministic file reading.

Fixtures are built in memory, never loaded from sample files, since those are generated in a
later phase and this suite must not depend on them.
"""

from __future__ import annotations

import io

from docx import Document
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject

from config import CONFIG
from rules.parse import ParsedDocument, UnreadableDocument, load_document

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _text_pdf_bytes(text: str, page_count: int = 1) -> bytes:
    """A PDF with a real content stream, built with pypdf's low level object API.

    pypdf has no text drawing helper of its own, so the page content stream is written by
    hand: a `Tj` show-text operator against the built in Helvetica font, which needs no font
    file to be embedded.
    """
    writer = PdfWriter()

    font = DictionaryObject()
    font[NameObject("/Type")] = NameObject("/Font")
    font[NameObject("/Subtype")] = NameObject("/Type1")
    font[NameObject("/BaseFont")] = NameObject("/Helvetica")
    font_ref = writer._add_object(font)

    for _ in range(page_count):
        page = writer.add_blank_page(width=300, height=300)

        stream = StreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 250 Td ({text}) Tj ET".encode())
        stream_ref = writer._add_object(stream)

        resources = DictionaryObject()
        font_dict = DictionaryObject()
        font_dict[NameObject("/F1")] = font_ref
        resources[NameObject("/Font")] = font_dict

        page[NameObject("/Contents")] = stream_ref
        page[NameObject("/Resources")] = resources

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _blank_pdf_bytes(page_count: int = 1) -> bytes:
    """A PDF with pages but no content stream at all, so `extract_text` returns nothing."""
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=200, height=200)

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _encrypted_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    # A distinct user password, not just an owner password, so an empty-password decrypt
    # attempt genuinely fails rather than silently succeeding.
    writer.encrypt(user_password="correct-horse", owner_password="battery-staple")

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx_bytes(paragraph_text: str, table_cell_text: str) -> bytes:
    document = Document()
    document.add_paragraph(paragraph_text)

    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = table_cell_text

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _png_bytes(size: tuple[int, int] = (50, 50)) -> bytes:
    image = Image.new("RGB", size, color="red")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_text_pdf_round_trips() -> None:
    # Protects against the extraction path silently dropping to the image branch for a
    # perfectly normal, machine-generated PDF.
    long_text = (
        "Hello, this is a real text layer with enough characters to clear "
        "the scanned threshold. "
    ) * 3
    data = _text_pdf_bytes(long_text)

    result = load_document(data, PDF_CONTENT_TYPE, "policy.pdf")

    assert isinstance(result, ParsedDocument)
    assert "Hello" in result.text
    assert result.image_only is False
    assert result.page_images == []


def test_scanned_pdf_is_treated_as_image_only() -> None:
    # Protects against a scanned PDF (or one with an unreadable text layer) being handed to
    # the model as empty text instead of being rendered to images.
    data = _blank_pdf_bytes(page_count=1)

    result = load_document(data, PDF_CONTENT_TYPE, "scan.pdf")

    assert result.image_only is True
    assert result.text == ""
    assert len(result.page_images) == 1
    assert result.page_images[0].startswith(b"\x89PNG")


def test_encrypted_pdf_raises_unreadable() -> None:
    # Protects against silently returning empty text for a file we could not actually open,
    # which would look like a successful, if content-free, parse.
    data = _encrypted_pdf_bytes()

    try:
        load_document(data, PDF_CONTENT_TYPE, "locked.pdf")
        raise AssertionError("expected UnreadableDocument")
    except UnreadableDocument:
        pass


def test_docx_round_trips_including_table_cell_text() -> None:
    # Protects against reading only `document.paragraphs`, which skips table content, and
    # claim forms are frequently laid out as tables.
    data = _docx_bytes("A paragraph with visible text.", "Value only present in a table cell")

    result = load_document(data, DOCX_CONTENT_TYPE, "claim.docx")

    assert "A paragraph with visible text." in result.text
    assert "Value only present in a table cell" in result.text
    assert result.image_only is False
    assert result.page_images == []


def test_png_gives_single_image_only_page() -> None:
    # Protects against an image being run through the text path, or split into more than
    # one page image.
    data = _png_bytes()

    result = load_document(data, "image/png", "photo.png")

    assert result.image_only is True
    assert result.text == ""
    assert len(result.page_images) == 1


def test_unsupported_content_type_raises_unreadable() -> None:
    # Protects against an unknown format being silently passed through as empty text
    # instead of failing loudly so the caller can report it.
    try:
        load_document(b"just some bytes", "text/plain", "notes.txt")
        raise AssertionError("expected UnreadableDocument")
    except UnreadableDocument:
        pass


def test_page_cap_truncates_and_records_observation() -> None:
    # Protects against a long scanned document blowing past the model's page limit; the cap
    # has to be enforced here, before the model call, not discovered during it.
    page_count = CONFIG.max_pages + 3
    data = _blank_pdf_bytes(page_count=page_count)

    result = load_document(data, PDF_CONTENT_TYPE, "long-scan.pdf")

    assert len(result.page_images) == CONFIG.max_pages
    assert any("dropped" in observation for observation in result.observations)


def test_large_image_is_downscaled_to_the_configured_edge() -> None:
    # Protects against an oversized image being sent to the model as is and rejected or
    # truncated by the provider's own per-image size limit.
    oversized = CONFIG.max_image_edge_px + 500
    data = _png_bytes(size=(oversized, 100))

    result = load_document(data, "image/png", "big.png")

    downscaled = Image.open(io.BytesIO(result.page_images[0]))
    assert max(downscaled.size) <= CONFIG.max_image_edge_px


def test_generic_content_type_falls_back_to_docx_extension() -> None:
    # Protects against browsers that send application/octet-stream (or nothing) for a DOCX
    # upload, which is common enough that the architecture calls it out explicitly.
    data = _docx_bytes("Fallback paragraph.", "Fallback table cell.")

    result = load_document(data, "application/octet-stream", "application-form.docx")

    assert "Fallback paragraph." in result.text
    assert "Fallback table cell." in result.text
    assert result.image_only is False
