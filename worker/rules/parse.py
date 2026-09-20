"""Deterministic file reading: the first node in the graph, and the only place that touches
raw file bytes directly.

Architecture section 5: PDF text via `pypdf`, DOCX via `python-docx`. A PDF with almost no
text is treated as scanned and rendered to images with `pypdfium2`. PyMuPDF is deliberately
avoided, it is AGPL licensed. Two limits are enforced here, before the model call, rather than
discovered during it: a page cap and a per-image size cap.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from PIL import Image
from pypdf import PdfReader

from config import CONFIG

# Content types recognised without looking at the filename. `image/*` covers every image
# subtype in one check rather than an enumeration that inevitably misses one.
_DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Filename fallback, used only when the content type is missing or generic. Browsers are not
# reliable about content type, so the extension is the second signal, not the first.
_EXTENSION_KIND = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".bmp": "image",
    ".tif": "image",
    ".tiff": "image",
    ".webp": "image",
}

# Generic content types a browser sends when it does not actually know the format.
_GENERIC_CONTENT_TYPES = {"", "application/octet-stream"}

_IMAGE_ONLY_OBSERVATION = (
    "Document is image only; snippet verification could not run against it."
)


class UnreadableDocument(Exception):
    """The file cannot be read at all: encrypted, corrupt, or an unsupported format."""


@dataclass(frozen=True)
class ParsedDocument:
    """The output of `load_document`, ready to hand to classification."""

    text: str
    page_images: list[bytes]
    image_only: bool
    observations: list[str]


def load_document(data: bytes, content_type: str, filename: str = "") -> ParsedDocument:
    """Read a file's bytes into text or page images, whichever the format actually has.

    Detection is content-type first, filename extension second, because a browser's content
    type is the more trustworthy signal when it is present and specific.
    """
    kind = _detect_kind(content_type, filename)

    if kind == "pdf":
        return _load_pdf(data)
    if kind == "docx":
        return _load_docx(data)
    if kind == "image":
        return _load_image(data)

    raise UnreadableDocument(f"Unsupported content type: {content_type!r}")


def _detect_kind(content_type: str, filename: str) -> str | None:
    normalised = (content_type or "").split(";", 1)[0].strip().lower()

    if normalised == "application/pdf":
        return "pdf"
    if normalised == _DOCX_CONTENT_TYPE:
        return "docx"
    if normalised.startswith("image/"):
        return "image"
    if normalised not in _GENERIC_CONTENT_TYPES:
        # A specific but unrecognised content type is a definite answer, not a reason to
        # fall back to the filename.
        return None

    return _EXTENSION_KIND.get(Path(filename).suffix.lower())


def _load_pdf(data: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        # Only an empty password is tried. A PDF still locked after that needs a password we
        # do not have, which is unreadable for our purposes either way.
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableDocument("PDF is encrypted and could not be opened")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except UnreadableDocument:
        raise
    except Exception as exc:
        raise UnreadableDocument(f"PDF could not be read: {exc}") from exc

    if len(text) >= CONFIG.scanned_text_threshold:
        return ParsedDocument(text=text, page_images=[], image_only=False, observations=[])

    # No usable text layer: treat it as scanned and fall back to page images.
    images, observations = _render_pdf_pages(data)
    observations.append(_IMAGE_ONLY_OBSERVATION)
    return ParsedDocument(text="", page_images=images, image_only=True, observations=observations)


def _render_pdf_pages(data: bytes) -> tuple[list[bytes], list[str]]:
    observations: list[str] = []
    try:
        pdf = pdfium.PdfDocument(data)
        total_pages = len(pdf)
        page_count = min(total_pages, CONFIG.max_pages)

        if total_pages > CONFIG.max_pages:
            dropped = total_pages - CONFIG.max_pages
            observations.append(
                f"Document has {total_pages} pages; only the first {CONFIG.max_pages} were "
                f"processed, {dropped} dropped."
            )

        images = [
            _encode_png(_downscale(pdf[index].render(scale=2.0).to_pil()))
            for index in range(page_count)
        ]
    except Exception as exc:
        raise UnreadableDocument(f"PDF could not be rendered: {exc}") from exc

    return images, observations


def _load_docx(data: bytes) -> ParsedDocument:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise UnreadableDocument(f"DOCX could not be read: {exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

    # `document.paragraphs` skips table content entirely, and forms are frequently laid out
    # as tables, so cell text is walked separately rather than assumed to be covered above.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return ParsedDocument(text="\n".join(parts), page_images=[], image_only=False, observations=[])


def _load_image(data: bytes) -> ParsedDocument:
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
    except Exception as exc:
        raise UnreadableDocument(f"Image could not be read: {exc}") from exc

    png_bytes = _encode_png(_downscale(image))
    return ParsedDocument(
        text="", page_images=[png_bytes], image_only=True, observations=[_IMAGE_ONLY_OBSERVATION]
    )


def _downscale(image: Image.Image) -> Image.Image:
    """Shrink an image so its longest edge fits the model's per-image limit.

    Never upscales: an image already under the limit is returned unchanged.
    """
    longest_edge = max(image.size)
    if longest_edge <= CONFIG.max_image_edge_px:
        return image

    scale = CONFIG.max_image_edge_px / longest_edge
    new_size = (max(1, round(image.width * scale)), max(1, round(image.height * scale)))
    return image.resize(new_size, Image.LANCZOS)


def _encode_png(image: Image.Image) -> bytes:
    # PNG has no direct encoder for every PIL mode a scanned page or a supplied image might
    # come in (CMYK, palette with transparency quirks, and so on), so anything unusual is
    # normalised to RGB first.
    if image.mode not in ("RGB", "RGBA", "L"):
        image = image.convert("RGB")

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
